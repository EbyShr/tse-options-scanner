"""
ماژول صدور گزارش اکسل چندشیت با قالب‌بندی حرفه‌ای، استایل‌های راست‌چین (RTL) و هشدارهای معاملاتی
"""

import os
import logging
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from analytics.underlying_momentum import determine_momentum_label

logger = logging.getLogger("OptionScanner.ExcelExporter")

# متن هشدار محدودیت‌های مدل بلک-شولز در بورس ایران
DISCLAIMER_NOTE = (
    "⚠ تذکر مهم درباره ارزش‌گذاری BSM در بازار ایران: "
    "نرخ بهره بدون ریسک و نوسان‌پذیری ضمنی (IV) در بازار اختیار معامله بورس تهران پروکسی دقیق و پیوسته‌ای ندارند. "
    "قیمت تئوریک بلک-شولز و عدد «حباب» صرفاً یک تخمین مقایسه‌ای نسبی جهت غربالگری میان نمادهاست، نه ارزش‌گذاری مطلق و قطعی. "
    "این سامانه صرفاً ابزار تحلیل است و هیچ‌گونه معامله‌ای اجرا نمی‌کند."
)


class ExcelExporter:
    """کلاس تولید گزارش‌های چندشیت اکسل با استایل کامل RTL و رنگ‌بندی تفکیک‌شده"""

    def __init__(self):
        # تعریف فونت‌ها و رنگ‌ها
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.title_font = Font(name="Calibri", size=13, bold=True, color="1B365D")
        self.disclaimer_font = Font(name="Calibri", size=10, italic=True, bold=True, color="8A1C14")
        self.regular_font = Font(name="Calibri", size=10)
        self.bold_font = Font(name="Calibri", size=10, bold=True)

        self.header_fill_main = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # سرمه‌ای
        self.header_fill_warn = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")  # قرمز
        self.header_fill_summary = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")  # آبی روشن
        self.banner_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # گلبهی ملایم

        # رنگ‌های وضعیت
        self.fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        self.fill_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        self.fill_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        self.thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

    def export(
        self,
        df_all: pd.DataFrame,
        underlying_stats: Dict[str, Dict[str, Any]],
        output_filepath: str,
    ) -> str:
        """تولید فایل اکسل کامل با ۳ شیت اصلی، شیت هشدار کم‌عمقی و شیت خلاصه دارایی‌های پایه"""
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        logger.info(f"در حال تولید گزارش اکسل در مسیر: {output_filepath}")

        wb = openpyxl.Workbook()
        # حذف شیت پیش‌فرض
        wb.remove(wb.active)

        # ۱. شیت اول: تمام قراردادها رتبه‌بندی شده
        ws_main = wb.create_sheet(title="تمام قراردادها (رتبه‌بندی)")
        self._populate_options_sheet(ws_main, df_all, is_warning_sheet=False)

        # ۲. شیت دوم: قراردادهای دارای هشدار نقدشوندگی (کم‌عمق)
        ws_warn = wb.create_sheet(title="قراردادهای کم‌عمق (هشدار)")
        df_illiquid = df_all[df_all["کم‌عمق"] == True].copy() if not df_all.empty else pd.DataFrame()
        self._populate_options_sheet(ws_warn, df_illiquid, is_warning_sheet=True)

        # ۳. شیت سوم: خلاصه دارایی‌های پایه
        ws_summary = wb.create_sheet(title="خلاصه دارایی‌های پایه")
        self._populate_underlying_summary_sheet(ws_summary, underlying_stats, df_all)

        wb.save(output_filepath)
        logger.info(f"فایل اکسل با موفقیت ذخیره شد: {output_filepath}")
        return output_filepath

    def _populate_options_sheet(self, ws, df: pd.DataFrame, is_warning_sheet: bool):
        """تنظیم و پر کردن شیت قراردادها همراه با بنر اخطار، هدر و فرمت‌بندی ستون‌ها"""
        ws.sheet_view.rightToLeft = True

        # سطر ۱: عنوان گزارش
        title = "گزارش جامع غربالگری و رتبه‌بندی اختیار معامله بورس تهران (TSETMC)" if not is_warning_sheet else "فهرست قراردادهای دارای ریسک نقدشوندگی بالا (کم‌عمق / کم‌معامله)"
        ws.cell(row=1, column=1, value=title).font = self.title_font
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=15)

        # سطر ۲: بنر سلب مسئولیت و توضیح محدودیت BSM
        ws.cell(row=2, column=1, value=DISCLAIMER_NOTE).font = self.disclaimer_font
        ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 36
        for col_idx in range(1, 16):
            ws.cell(row=2, column=col_idx).fill = self.banner_fill
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=15)

        if df.empty:
            ws.cell(row=4, column=1, value="هیچ قراردادی در این دسته یافت نشد.").font = self.regular_font
            return

        # ستون‌های نمایشی انتخابی و مرتب‌شده برای کاربر
        display_cols = [
            "امتیاز ترکیبی",
            "شرح فرمول امتیاز",
            "نماد",
            "دارایی پایه",
            "نوع قرارداد",
            "قیمت اعمال",
            "قیمت پایه",
            "قیمت پایانی بازار",
            "قیمت تئوریک BSM",
            "حباب خام (%)",
            "جذابیت حباب (معکوس)",
            "اهرم",
            "رتبه صدکی اهرم",
            "وضعیت نقدینگی",
            "رتبه صدکی نقدینگی",
            "ارزش معاملات امروز (ریال)",
            "میانگین ارزش ۵ روزه (ریال)",
            "تعداد معاملات امروز",
            "نسبت معامله به میانگین",
            "وضعیت سودآوری",
            "فاصله تا سربه‌سر (%)",
            "روند دارایی پایه",
            "هم‌جهتی با روند",
            "روزهای تا سررسید (DTE)",
            "تاریخ سررسید",
            "دلتا",
            "نوسان‌پذیری",
            "بهترین مظنه خرید (Bid)",
            "بهترین مظنه فروش (Ask)",
            "اسپرد (اختلاف مظنه)",
        ]

        # فیلتر ستون‌های موجود در دیتافریم
        active_cols = [c for c in display_cols if c in df.columns]

        # سطر ۴: درج هدرها
        header_fill = self.header_fill_warn if is_warning_sheet else self.header_fill_main
        ws.row_dimensions[4].height = 26
        for c_idx, col_name in enumerate(active_cols, 1):
            cell = ws.cell(row=4, column=c_idx, value=col_name)
            cell.font = self.header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        # درج داده‌ها
        start_row = 5
        for r_idx, (_, row_data) in enumerate(df.iterrows(), start_row):
            ws.row_dimensions[r_idx].height = 20
            is_illiquid = bool(row_data.get("کم‌عمق", False))

            for c_idx, col_name in enumerate(active_cols, 1):
                val = row_data[col_name]
                cell = ws.cell(row=r_idx, column=c_idx)

                # تنظیم فونت و ترازبندی
                cell.font = self.regular_font
                cell.border = self.thin_border

                # فرمت‌بندی اعداد
                if isinstance(val, (int, float)):
                    if "ریال" in col_name or "قیمت" in col_name or "مظنه" in col_name or "اسپرد" in col_name:
                        cell.value = val
                        cell.number_format = "#,##0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif "%" in col_name or "رتبه" in col_name or "امتیاز" in col_name or "اهرم" in col_name or "دلتا" in col_name or "نسبت" in col_name:
                        cell.value = val
                        cell.number_format = "0.0" if abs(val) < 1000 else "#,##0.0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.value = val
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.value = str(val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # رنگ‌آمیزی ستون وضعیت نقدینگی
                if col_name == "وضعیت نقدینگی":
                    if is_illiquid:
                        cell.fill = self.fill_red
                        cell.font = self.bold_font
                    else:
                        cell.fill = self.fill_green

                # رنگ‌آمیزی ستون روند دارایی پایه
                if col_name == "روند دارایی پایه":
                    txt = str(val)
                    if "صعودی" in txt:
                        cell.fill = self.fill_green
                    elif "نزولی" in txt:
                        cell.fill = self.fill_red

                # رنگ‌آمیزی هم‌جهتی با روند
                if col_name == "هم‌جهتی با روند":
                    txt = str(val)
                    if "✓" in txt:
                        cell.fill = self.fill_green
                    elif "⚠" in txt:
                        cell.fill = self.fill_red

        # تنظیم اتوماتیک عرض ستون‌ها
        for c_idx, col_name in enumerate(active_cols, 1):
            col_letter = get_column_letter(c_idx)
            max_len = max(len(str(col_name)), 10)
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    def _populate_underlying_summary_sheet(
        self,
        ws,
        underlying_stats: Dict[str, Dict[str, Any]],
        df_all: pd.DataFrame,
    ):
        """تنظیم و پر کردن شیت خلاصه وضعیت دارایی‌های پایه"""
        ws.sheet_view.rightToLeft = True

        ws.cell(row=1, column=1, value="خلاصه آمار مومنتوم و روند دارایی‌های پایه هدف").font = self.title_font
        ws.row_dimensions[1].height = 25
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)

        headers = [
            "نماد پایه",
            "قیمت پایانی",
            "برچسب روند و مومنتوم",
            "وضعیت صف معاملاتی",
            "بازدهی ۱ روزه (%)",
            "بازدهی ۳ روزه (%)",
            "بازدهی ۵ روزه (%)",
            "بازدهی ۱۰ روزه (%)",
            "میانگین SMA5",
            "فاصله از SMA5 (%)",
            "میانگین SMA20",
            "فاصله از SMA20 (%)",
            "نوسان تاریخی سالانه (%)",
            "تعداد قرارداد فعال Call",
            "تعداد قرارداد فعال Put",
            "ارزش کل معاملات آپشن (ریال)",
        ]

        ws.row_dimensions[3].height = 26
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=c_idx, value=h)
            cell.font = self.header_font
            cell.fill = self.header_fill_summary
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.thin_border

        # تجمیع آمار آپشن‌ها به تفکیک پایه
        opt_counts: Dict[str, Dict[str, Any]] = {}
        if not df_all.empty:
            for _, r in df_all.iterrows():
                u_sym = str(r.get("دارایی پایه", ""))
                opt_counts.setdefault(u_sym, {"calls": 0, "puts": 0, "total_val": 0.0})
                is_call = "Call" in str(r.get("نوع قرارداد", ""))
                if is_call:
                    opt_counts[u_sym]["calls"] += 1
                else:
                    opt_counts[u_sym]["puts"] += 1
                opt_counts[u_sym]["total_val"] += float(r.get("ارزش معاملات امروز (ریال)", 0.0))

        start_row = 4
        # جلوگیری از تکرار نمادهای نرمال‌شده
        seen_syms = set()
        for sym, stats in underlying_stats.items():
            norm_s = stats.get("NormalizedSymbol", sym)
            if norm_s in seen_syms:
                continue
            seen_syms.add(norm_s)

            u_info = opt_counts.get(sym) or opt_counts.get(norm_s) or {"calls": 0, "puts": 0, "total_val": 0.0}
            trend_label = determine_momentum_label(stats)

            row_vals = [
                stats.get("Symbol", sym),
                stats.get("Close", 0.0),
                trend_label,
                stats.get("QueueStatus", "متعادل"),
                stats.get("Return1D", 0.0),
                stats.get("Return3D", 0.0),
                stats.get("Return5D", 0.0),
                stats.get("Return10D", 0.0),
                stats.get("SMA5", 0.0),
                stats.get("DistSMA5", 0.0),
                stats.get("SMA20", 0.0),
                stats.get("DistSMA20", 0.0),
                round(stats.get("RealizedVol", 0.35) * 100.0, 1),
                u_info["calls"],
                u_info["puts"],
                round(u_info["total_val"]),
            ]

            ws.row_dimensions[start_row].height = 20
            for c_idx, v in enumerate(row_vals, 1):
                cell = ws.cell(row=start_row, column=c_idx, value=v)
                cell.font = self.regular_font
                cell.border = self.thin_border

                if isinstance(v, (int, float)):
                    if c_idx in (2, 9, 11, 16):  # مبالغ
                        cell.number_format = "#,##0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif c_idx in (5, 6, 7, 8, 10, 12, 13):  # درصدها
                        cell.number_format = "0.0"
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # استایل روند
                if c_idx == 3:
                    if "صعودی" in str(v):
                        cell.fill = self.fill_green
                        cell.font = self.bold_font
                    elif "نزولی" in str(v):
                        cell.fill = self.fill_red
                        cell.font = self.bold_font

            start_row += 1

        for c_idx, h in enumerate(headers, 1):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = max(len(str(h)) + 4, 14)
